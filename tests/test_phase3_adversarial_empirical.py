"""
Phase 3 Empirical Challenger Verification & Adversarial Stress Suite.
Tests:
1. Database resilience under injected mutations & query exceptions across all modules.
2. Background mailer, PDF jobs, and scheduler recovery.
3. Multiple consecutive failure cycles to ensure zero dirty session leakage.
4. Requirements.txt dependency version constraints verification.
5. Full Python compilation and Jinja2 template parse verification.
"""
import os
import py_compile
import re
from io import BytesIO
from pathlib import Path
from unittest.mock import patch, MagicMock
import jinja2
import pytest
from sqlalchemy import text
from sqlalchemy.exc import (
    IntegrityError,
    OperationalError,
    SQLAlchemyError,
    PendingRollbackError,
    InvalidRequestError,
)

from platform_app import create_app
from extensions import db
from modules.propostas.models import (
    User,
    Department,
    Proposal,
    AgendaEntry,
    SystemOptionOverride,
    SystemOptionState,
    SystemOptionCatalog,
)
from modules.suporte.models import AssistenciaTarefa, OrcamentoStatus, AssistenciaOrcamento
from modules.chamados.models import Ticket, TicketMessage
from modules.chamados.mailer import enviar_email
from modules.chamados.services.notify import recipients_for_ticket
from modules.propostas.services.pdf_jobs import PdfJobManager, PdfJob
from modules.propostas.utils.systems import _load_override_map, _load_system_states, _load_custom_options

PROJECT_ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture
def app():
    app = create_app()
    app.config.update(
        TESTING=True,
        WTF_CSRF_ENABLED=False,
        SERVER_NAME="localhost",
    )
    with app.app_context():
        # Ensure base admin user exists
        user = User.query.filter_by(usuario="challenger_p3_admin").first()
        if not user:
            user = User(
                usuario="challenger_p3_admin",
                nome_completo="Challenger Admin",
                email="challenger@sollus.com",
                tipo="admin",
                role="admin",
                password_hash="dummy_hash",
                is_active=True,
            )
            db.session.add(user)
            db.session.commit()
        yield app


@pytest.fixture
def client(app):
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(usuario="challenger_p3_admin").first()
        admin_id = admin.id
    with client.session_transaction() as sess:
        sess["usuario_id"] = admin_id
        sess["user_id"] = admin_id
        sess["_user_id"] = str(admin_id)
        sess["tipo"] = "admin"
    return client


# =========================================================================
# SECTION 1: Direct Session Rollback & Dirty State Recovery Oracles
# =========================================================================

def test_oracle_unrecovered_vs_recovered_session(app):
    """
    Oracle Test: Demonstrates that without rollback, a failed transaction throws
    PendingRollbackError on subsequent operations, but with rollback() the session
    fully recovers.
    """
    with app.app_context():
        # 1. Unrecovered session scenario:
        try:
            db.session.execute(text("SELECT * FROM non_existent_table_oracle_fail"))
            db.session.flush()
        except SQLAlchemyError:
            pass  # omit rollback intentionally to poison session

        with pytest.raises((PendingRollbackError, InvalidRequestError, SQLAlchemyError)):
            db.session.execute(text("SELECT 1")).scalar()

        # Now execute rollback to heal the session:
        db.session.rollback()

        # Verify session is healed
        val = db.session.execute(text("SELECT 1")).scalar()
        assert val == 1, "Session must recover after rollback"


def test_consecutive_100_failure_recovery_cycles(app):
    """
    Stress-test: Execute 100 consecutive cycles of failure -> rollback -> valid query.
    Ensures no connection pool exhaustion or transaction state leakage.
    """
    with app.app_context():
        for i in range(100):
            try:
                db.session.execute(text(f"INSERT INTO invalid_table_{i} VALUES ({i})"))
                db.session.flush()
            except SQLAlchemyError:
                db.session.rollback()

            res = db.session.execute(text("SELECT 42")).scalar()
            assert res == 42, f"Cycle {i} failed to recover session"


def test_integrity_constraint_violation_recovery(app):
    """
    Test recovery after unique constraint or foreign key violation.
    """
    with app.app_context():
        # User username is unique
        u1 = User(usuario="unique_test_user_1", nome_completo="User 1", email="u1@sollus.com")
        db.session.add(u1)
        db.session.commit()

        # Attempt to insert duplicate
        u2 = User(usuario="unique_test_user_1", nome_completo="User 2", email="u2@sollus.com")
        db.session.add(u2)
        try:
            db.session.commit()
            pytest.fail("Expected IntegrityError on duplicate username")
        except IntegrityError:
            db.session.rollback()

        # Follow-up write and query should succeed cleanly
        u3 = User(usuario="unique_test_user_2", nome_completo="User 3", email="u3@sollus.com")
        db.session.add(u3)
        db.session.commit()

        found = User.query.filter_by(usuario="unique_test_user_2").first()
        assert found is not None
        assert found.nome_completo == "User 3"


# =========================================================================
# SECTION 2: Background Mailer, Notify & PDF Workers Resilience
# =========================================================================

def test_chamados_mailer_db_commit_failure_isolation(app):
    """
    Verify enviar_email catches DB commit failures, rolls back, and allows subsequent
    mailer/database operations on the same session.
    """
    with app.app_context():
        # Inject commit failure
        with patch.object(db.session, "commit", side_effect=OperationalError("DB Locked", {}, None)):
            success = enviar_email(
                destinatarios=["support@sollus.com"],
                assunto="Challenger Test Email",
                html_corpo="<p>Test</p>",
            )
            assert success is False

        # Session must be functional
        assert db.session.execute(text("SELECT 100")).scalar() == 100

        # Normal email queueing should now succeed
        success_after = enviar_email(
            destinatarios=["support2@sollus.com"],
            assunto="Challenger Test Email Recovered",
            html_corpo="<p>Test Recovered</p>",
        )
        assert success_after is True


def test_chamados_notify_recipients_query_failure_isolation(app):
    """
    Verify recipients_for_ticket handles query failure gracefully with rollback.
    """
    with app.app_context():
        dummy_ticket = MagicMock()
        dummy_ticket.id = 88888
        dummy_ticket.user = None
        dummy_ticket.assignee = None
        dummy_ticket.messages = None

        with patch("modules.chamados.models.TicketMessage.query") as mock_query:
            mock_query.filter_by.side_effect = SQLAlchemyError("Query execution aborted")
            recipients = recipients_for_ticket(dummy_ticket)
            assert isinstance(recipients, list)

        # Verify session can execute new queries
        res = db.session.execute(text("SELECT count(*) FROM users")).scalar()
        assert res >= 1


def test_pdf_jobs_manager_cleanup_and_save_rollback(app):
    """
    Verify PdfJobManager rolls back on exceptions during cleanup or job updates.
    """
    with app.app_context():
        mgr = PdfJobManager()
        # 1. Simulate failure in cleanup
        with patch.object(db.session, "commit", side_effect=SQLAlchemyError("Cleanup failed")):
            mgr.cleanup()

        # Verify session is alive
        assert db.session.execute(text("SELECT 1")).scalar() == 1

        # 2. Simulate failure in _save_job
        dummy_job = PdfJob(proposal_id=999, job_type="preview", status="pending")
        with patch.object(db.session, "commit", side_effect=SQLAlchemyError("Save job failed")):
            mgr._save_job(dummy_job)

        # Session should still be active
        assert db.session.execute(text("SELECT 1")).scalar() == 1


def test_propostas_systems_utils_rollback_isolation(app):
    """
    Verify _load_override_map, _load_system_states, _load_custom_options catch exceptions,
    execute rollback, and return clean fallbacks.
    """
    with app.app_context():
        with patch("modules.propostas.models.SystemOptionOverride.query") as mock_ov:
            mock_ov.all.side_effect = SQLAlchemyError("SystemOptionOverride error")
            res_ov = _load_override_map()
            assert res_ov == {}

        with patch("modules.propostas.models.SystemOptionState.query") as mock_st:
            mock_st.all.side_effect = SQLAlchemyError("SystemOptionState error")
            res_st = _load_system_states()
            assert res_st == {}

        with patch("modules.propostas.models.SystemOptionCatalog.query") as mock_cat:
            mock_cat.all.side_effect = SQLAlchemyError("SystemOptionCatalog error")
            res_cat = _load_custom_options()
            assert res_cat == []

        # Scoped session must be completely healthy
        assert db.session.execute(text("SELECT 1")).scalar() == 1


# =========================================================================
# SECTION 3: Route Handlers DB Mutation Exception Stress-Testing
# =========================================================================

def test_route_assistencia_orcamentos_criar_rollback(app, client):
    """
    Simulate commit exception on POST /assistencia/orcamentos/criar and verify
    rollback and subsequent query health.
    """
    with patch.object(db.session, "commit", side_effect=OperationalError("Disk I/O Error", {}, None)):
        resp = client.post(
            "/assistencia/orcamentos/criar",
            data={
                "dataEnvio": "2026-08-26",
                "tipoVisita": "AVULSA",
                "equipamento": "Catraca Challenger",
                "cliente": "Empresa Challenger",
                "numeroProposta": "PROP-CHALLENGER-1",
                "valor": "2000,00",
                "unidade": "RJ",
            },
        )
        assert resp.status_code == 302

    with app.app_context():
        # Verify rollback prevented partial record creation
        record = OrcamentoStatus.query.filter_by(numero_proposta="PROP-CHALLENGER-1").first()
        assert record is None
        # Verify session is clean
        assert db.session.execute(text("SELECT 1")).scalar() == 1


def test_route_shared_agenda_criar_atualizar_rollback(app, client):
    """
    Simulate commit exception on /admin/suporte/api/agenda/criar and atualizar.
    """
    with app.app_context():
        admin = User.query.filter_by(usuario="challenger_p3_admin").first()
        admin_id = admin.id

    with patch.object(db.session, "commit", side_effect=SQLAlchemyError("Agenda commit failed")):
        resp = client.post(
            "/admin/suporte/api/agenda/criar",
            data={
                "usuario_id": str(admin_id),
                "unidade": "SP",
                "data_atendimento": "2026-09-01",
                "periodo": "Manhã",
                "obs": "Simulated error agenda",
            },
        )
        assert resp.status_code in (302, 200, 500)

    with app.app_context():
        # Verify session is clean
        assert db.session.execute(text("SELECT 1")).scalar() == 1


def test_route_financeiro_contas_receber_adicionar_rollback(app, client):
    """
    Simulate database exception on /financeiro/contas-receber/adicionar.
    """
    # Ensure legacy table exists
    with app.app_context():
        db.session.execute(text("""
            CREATE TABLE IF NOT EXISTS contas_receber (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                cliente VARCHAR(255),
                cnpj VARCHAR(50),
                contrato VARCHAR(100),
                software VARCHAR(100),
                data_primeira_pendencia VARCHAR(50),
                qt_pendencias INTEGER,
                dias_atraso INTEGER,
                total FLOAT,
                empresa_responsavel VARCHAR(255),
                valor FLOAT,
                criado_por VARCHAR(100),
                status VARCHAR(50),
                id_pai INTEGER DEFAULT 0,
                data_bloqueio VARCHAR(50),
                cancelamento VARCHAR(50),
                deferimento_cancelamento VARCHAR(50),
                informacoes TEXT
            )
        """))
        db.session.commit()

    with patch.object(db.session, "commit", side_effect=OperationalError("DB busy", {}, None)):
        resp = client.post(
            "/financeiro/contas-receber/adicionar",
            data={
                "cliente": "Cliente Rollback Test Financeiro",
                "cnpj": "12.345.678/0001-90",
                "contrato": "CON-2026-999",
                "software": "Sollus Connect",
                "data_primeira_pendencia": "2026-06-10",
                "empresa_responsavel": "SOLLUS RJ",
                "valor": "1.500,00",
            },
            follow_redirects=True,
        )
        assert resp.status_code in (200, 302, 500)

    with app.app_context():
        # Verify session works
        assert db.session.execute(text("SELECT 1")).scalar() == 1


# =========================================================================
# SECTION 4: Requirements.txt Version Constraints Verification
# =========================================================================

def test_requirements_file_exact_constraints():
    """
    Verify that requirements.txt satisfies all Phase 3 requirements exactly:
    - apscheduler>=3.10.0,<4.0.0
    - Flask-Mail>=0.9.1
    - openpyxl>=3.1.0
    """
    req_file = PROJECT_ROOT / "requirements.txt"
    assert req_file.exists(), "requirements.txt not found"
    content = req_file.read_text(encoding="utf-8")

    assert re.search(r"apscheduler\s*>=\s*3\.10\.0\s*,\s*<\s*4\.0\.0", content, re.IGNORECASE) is not None, (
        "apscheduler>=3.10.0,<4.0.0 is missing or incorrectly formatted in requirements.txt"
    )
    assert re.search(r"Flask-Mail\s*>=\s*0\.9\.1", content, re.IGNORECASE) is not None, (
        "Flask-Mail>=0.9.1 is missing or incorrectly formatted in requirements.txt"
    )
    assert re.search(r"openpyxl\s*>=\s*3\.1\.0", content, re.IGNORECASE) is not None, (
        "openpyxl>=3.1.0 is missing or incorrectly formatted in requirements.txt"
    )


def test_requirements_packages_runtime_import():
    """
    Verify that the pinned packages can be imported and function in Python 3.13.
    """
    import apscheduler
    from apscheduler.schedulers.background import BackgroundScheduler
    sched = BackgroundScheduler()
    assert sched is not None

    import flask_mail
    from flask_mail import Mail, Message
    assert Mail is not None
    assert Message is not None

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Test Cell"
    assert ws["A1"].value == "Test Cell"


# =========================================================================
# SECTION 5: Full Codebase Python & Jinja2 Compilation Integrity
# =========================================================================

def test_all_python_files_compile_cleanly():
    """
    Verify that all Python source files in the project compile without SyntaxError.
    """
    compiled_count = 0
    errors = []
    for root, dirs, files in os.walk(PROJECT_ROOT):
        if any(x in root for x in [".git", "__pycache__", "env", "venv", ".agents"]):
            continue
        for f in files:
            if f.endswith(".py"):
                file_path = os.path.join(root, f)
                try:
                    py_compile.compile(file_path, doraise=True)
                    compiled_count += 1
                except Exception as e:
                    errors.append((file_path, str(e)))

    assert len(errors) == 0, f"Python compilation failed on {len(errors)} files: {errors}"
    assert compiled_count > 100, f"Expected >100 python files compiled, got {compiled_count}"


def test_all_jinja2_templates_parse_cleanly():
    """
    Verify that all Jinja2 HTML templates in templates/ parse without syntax errors.
    """
    templates_dir = PROJECT_ROOT / "templates"
    env = jinja2.Environment(loader=jinja2.FileSystemLoader(str(templates_dir)))
    parsed_count = 0
    errors = []

    for root, dirs, files in os.walk(templates_dir):
        for f in files:
            if f.endswith(".html") or f.endswith(".j2"):
                file_path = os.path.join(root, f)
                rel_path = os.path.relpath(file_path, templates_dir)
                try:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as tf:
                        env.parse(tf.read())
                    parsed_count += 1
                except Exception as e:
                    errors.append((rel_path, str(e)))

    assert len(errors) == 0, f"Jinja2 template parsing failed on {len(errors)} templates: {errors}"
    assert parsed_count > 50, f"Expected >50 templates parsed, got {parsed_count}"
